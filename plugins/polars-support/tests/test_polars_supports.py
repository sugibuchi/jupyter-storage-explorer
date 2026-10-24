"""polars-support: its supports' views (polars reads by URL: local files) and
code, its options, and its entry point."""

from __future__ import annotations

import io
from importlib.metadata import entry_points

import openpyxl
import polars as pl
import pytest
from jupyter_storage_explorer_polars_support import (
    CsvByPolars,
    CsvByPolarsOptions,
    DeltaLakeByPolars,
    DeltaLakeByPolarsOptions,
    ExcelByPolars,
    ExcelByPolarsOptions,
    JsonByPolars,
    JsonByPolarsOptions,
    JsonLinesByPolars,
    ParquetByPolars,
    ParquetByPolarsOptions,
    PolarsOptions,
    PolarsSupport,
)
from jupyter_storage_explorer_polars_support.supports import reads_itself

from jupyter_storage_explorer.filesystems import S3Location
from jupyter_storage_explorer.formats import (
    Csv,
    DeltaLake,
    Excel,
    Json,
    JsonLines,
    Parquet,
    QueryView,
    TableSetView,
    ViewLimits,
    ViewWithChoice,
)
from jupyter_storage_explorer.formats.packages import PyArrow
from jupyter_storage_explorer.testing import local_item

LIMITS = ViewLimits(max_display_rows=2)


def parquet_bytes(df: pl.DataFrame) -> bytes:
    buffer = io.BytesIO()
    df.write_parquet(buffer)
    return buffer.getvalue()


def workbook() -> bytes:
    """Two sheets: "one", a column "a" of 1, 2, 3; "two", a column "b"."""
    book = openpyxl.Workbook()
    book.active.title = "one"
    for row in (["a"], [1], [2], [3]):
        book.active.append(row)
    book.create_sheet("two").append(["b"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def on_s3(key: str) -> S3Location:
    """A location polars reads itself: its code reads the URL."""
    return S3Location(bucket="b", key=key)


def last_line(code: str) -> str:
    return code.strip().splitlines()[-1]


def data(support, fmt, item, options=None):
    """The data of the support's view of ``item``: a QueryView's."""
    view = support.view(fmt, item, options or support.options_model(), LIMITS)
    assert isinstance(view, QueryView)
    return view.data


def test_plugin_registered():
    [entry_point] = entry_points(
        group="jupyter_storage_explorer.plugins", name="polars-support"
    )
    plugin = entry_point.load()()
    assert isinstance(plugin, PolarsSupport)
    assert plugin.description == "Polars support for Storage Explorer"
    assert CsvByPolars in plugin.format_supports


class TestViews:
    def test_csv_with_its_separator(self, tmp_path):
        path = tmp_path / "a.csv"
        path.write_text("x;y\n1;2\n3;4\n")
        df = data(CsvByPolars(), Csv(), local_item(path)).collect()
        assert df.columns == ["x", "y"]  # the separator detected
        assert df.rows() == [(1, 2), (3, 4)]

    @pytest.mark.parametrize(
        ("support", "fmt", "write", "name"),
        [
            (
                ParquetByPolars(),
                Parquet(),
                lambda df, p: df.write_parquet(p),
                "a.parquet",
            ),
            (CsvByPolars(), Csv(), lambda df, p: df.write_csv(p), "a.csv"),
            (
                JsonLinesByPolars(),
                JsonLines(),
                lambda df, p: df.write_ndjson(p),
                "a.jsonl",
            ),
        ],
    )
    def test_local_file(self, tmp_path, support, fmt, write, name):
        path = tmp_path / name
        write(pl.DataFrame({"a": [1, 2]}), path)
        assert data(support, fmt, local_item(path)).collect()["a"].to_list() == [1, 2]

    def test_delta_table(self, tmp_path):
        pl.DataFrame({"a": [1, 2]}).write_delta(tmp_path / "t")
        frame = data(DeltaLakeByPolars(), DeltaLake(), local_item(tmp_path / "t"))
        assert frame.collect().height == 2

    def test_excel_sheets_on_tabs(self, make_item):
        # polars reads the workbook through the explorer's filesystem: each
        # sheet whole (a DataFrame), queried in memory, when chosen.
        item = make_item("/a.xlsx", workbook())
        view = ExcelByPolars().view(Excel(), item, ExcelByPolarsOptions(), LIMITS)
        assert isinstance(view, TableSetView)
        assert view.selector == "tab"
        assert [(n.label, n.id) for n in view.nodes] == [("one", "one"), ("two", "two")]
        assert view.loader("one")["a"].to_list() == [1, 2, 3]
        assert view.loader("two").columns == ["b"]

    def test_excel_with_one_sheet(self, make_item):
        book = openpyxl.Workbook()
        book.active.append(["a"])
        book.active.append([1])
        buffer = io.BytesIO()
        book.save(buffer)
        item = make_item("/a.xlsx", buffer.getvalue())
        frame = data(ExcelByPolars(), Excel(), item, ExcelByPolarsOptions())
        assert frame["a"].to_list() == [1]

    def test_parquet_folder_skips_hidden_files(self, tmp_path):
        pl.DataFrame({"a": [1]}).write_parquet(tmp_path / "part-0.parquet")
        pl.DataFrame({"a": [2]}).write_parquet(tmp_path / "part-1.parquet")
        (tmp_path / "_SUCCESS").write_bytes(b"")
        (tmp_path / "sub").mkdir()
        pl.DataFrame({"a": [3]}).write_parquet(tmp_path / "sub" / "part-0.parquet")
        df = data(ParquetByPolars(), Parquet(), local_item(tmp_path)).collect()
        assert sorted(df["a"].to_list()) == [1, 2]


class TestPartitionedParquet:
    def write(self, root):
        for year in (2025, 2026):
            for month in (1, 2):
                folder = root / f"year={year}" / f"month={month}"
                folder.mkdir(parents=True)
                pl.DataFrame({"a": [year * 100 + month]}).write_parquet(
                    folder / "part-0.parquet"
                )
                (folder / "_SUCCESS").write_bytes(b"")
        # Left by a writer: not part of the table.
        (root / "_temporary" / "0").mkdir(parents=True)
        pl.DataFrame({"a": [-1]}).write_parquet(root / "_temporary" / "0" / "x.parquet")

    def test_scan_with_partition_columns(self, tmp_path):
        self.write(tmp_path)
        fmt = Parquet()
        assert fmt.matches(local_item(tmp_path))
        df = data(ParquetByPolars(), fmt, local_item(tmp_path)).collect().sort("a")
        assert df.columns == ["a", "year", "month"]
        assert df.rows() == [
            (202501, 2025, 1),
            (202502, 2025, 2),
            (202601, 2026, 1),
            (202602, 2026, 2),
        ]

    def test_scan_of_a_folder_not_detected(self, tmp_path):
        # Its layout is found then, through the filesystem given.
        self.write(tmp_path)
        df = data(ParquetByPolars(), Parquet(), local_item(tmp_path)).collect()
        assert (df.columns, df.height) == (["a", "year", "month"], 4)

    def test_scan_hive_files(self, tmp_path):
        folder = tmp_path / "dt=2026-10-01"
        folder.mkdir()
        pl.DataFrame({"a": [1, 2]}).write_parquet(folder / "000000_0")
        (folder / "_SUCCESS").write_bytes(b"")
        fmt = Parquet()
        assert fmt.matches(local_item(tmp_path))
        df = data(ParquetByPolars(), fmt, local_item(tmp_path)).collect()
        assert df["a"].to_list() == [1, 2]
        assert df.columns == ["a", "dt"]

    def test_code(self, make_item, memfs):
        memfs.pipe_file("/t/year=2025/month=1/part-0.parquet", b"")
        code = ParquetByPolars().recipe(Parquet(), make_item("/t"), PolarsOptions())
        assert last_line(code["Load as LazyFrame"]) == (
            "df = pl.scan_parquet('memory:///t/year=*/month=*/*.parquet', "
            "hive_partitioning=True)"
        )
        assert last_line(code["Load as DataFrame"]) == (
            "df = pl.read_parquet('memory:///t/year=*/month=*/*.parquet', "
            "hive_partitioning=True)"
        )

    def test_code_for_hive_files(self, make_item, memfs):
        frame = parquet_bytes(pl.DataFrame({"a": [1]}))
        memfs.pipe_file("/t/dt=2026-10-01/000000_0", frame)
        code = ParquetByPolars().recipe(Parquet(), make_item("/t"), PolarsOptions())
        options = "'memory:///t/dt=*/*', hive_partitioning=True, hidden_file_prefix=['.', '_'])"
        assert last_line(code["Load as LazyFrame"]) == f"df = pl.scan_parquet({options}"
        # read_parquet takes no hidden_file_prefix: scanned, then read.
        assert last_line(code["Load as DataFrame"]) == (
            f"df = pl.scan_parquet({options}.collect()"
        )


class TestRecipes:
    def test_csv(self, make_item):
        code = CsvByPolars().recipe(
            Csv(),
            make_item("/a.tsv", b"x\ty\n", location=on_s3("a.tsv")),
            PolarsOptions(),
        )
        [frame] = code.values()
        assert frame.startswith("import polars as pl")
        assert last_line(frame) == "df = pl.read_csv('s3://b/a.tsv', separator='\\t')"

    def test_lazy_and_eager(self, make_item, memfs):
        memfs.pipe_file("/d/part-0.parquet", b"")
        code = ParquetByPolars().recipe(Parquet(), make_item("/d"), PolarsOptions())
        assert last_line(code["Load as LazyFrame"]) == (
            "df = pl.scan_parquet('memory:///d/*.parquet')"
        )
        assert last_line(code["Load as DataFrame"]) == (
            "df = pl.read_parquet('memory:///d/*.parquet')"
        )
        lines = JsonLinesByPolars().recipe(
            JsonLines(),
            make_item("/a.jsonl", b"{}\n", location=on_s3("a.jsonl")),
            PolarsOptions(),
        )
        assert last_line(lines["Load as LazyFrame"]) == (
            "df = pl.scan_ndjson('s3://b/a.jsonl')"
        )
        assert last_line(lines["Load as DataFrame"]) == (
            "df = pl.read_ndjson('s3://b/a.jsonl')"
        )

    def test_delta(self, make_item, memfs):
        memfs.pipe_file("/t/_delta_log/0.json", b"{}")
        item = make_item("/t", location=S3Location(bucket="b", key="t"))
        code = DeltaLakeByPolars().recipe(DeltaLake(), item, DeltaLakeByPolarsOptions())
        assert last_line(code["Load as LazyFrame"]) == "df = pl.scan_delta('s3://b/t')"
        assert last_line(code["Load as DataFrame"]) == "df = pl.read_delta('s3://b/t')"

    def test_json(self, make_item):
        code = JsonByPolars().recipe(
            Json(),
            make_item("/a.json", b"{}", location=on_s3("a.json")),
            PolarsOptions(),
        )
        assert last_line(code["Load as DataFrame"]) == (
            "df = pl.read_json('s3://b/a.json')"
        )


class TestOptions:
    def test_storage_options_without_secrets(self, make_item):
        options = PolarsOptions(
            storage_options={"aws_region": "eu", "aws_secret_access_key": "s"}
        )
        [code] = (
            CsvByPolars()
            .recipe(
                Csv(), make_item("/a.csv", b"x,y\n", location=on_s3("a.csv")), options
            )
            .values()
        )
        assert last_line(code) == (
            "df = pl.read_csv('s3://b/a.csv', storage_options={'aws_region': 'eu'})"
        )

    def test_delta_options(self, make_item, memfs):
        memfs.pipe_file("/t/_delta_log/0.json", b"{}")
        options = DeltaLakeByPolarsOptions(
            storage_options={"region": "eu"},
            use_pyarrow=True,
            delta_table_options={"version": 1},
        )
        item = make_item("/t", location=S3Location(bucket="b", key="t"))
        code = DeltaLakeByPolars().recipe(DeltaLake(), item, options)
        expected = (
            "'s3://b/t', delta_table_options={'version': 1}, use_pyarrow=True, "
            "storage_options={'region': 'eu'})"
        )
        assert last_line(code["Load as LazyFrame"]) == f"df = pl.scan_delta({expected}"
        assert last_line(code["Load as DataFrame"]) == f"df = pl.read_delta({expected}"

    def test_option_names(self):
        assert PolarsOptions.option_names() == {
            "storage_options",
            "credential_provider",
            "fsspec_storage_options",
        }
        assert DeltaLakeByPolarsOptions.option_names() >= {
            "use_pyarrow",
            "storage_options",
        }


class TestReaderOptions:
    """Each format's options are its reader's arguments, given or left to
    polars."""

    def test_csv_arguments_in_the_data_browser(self, tmp_path):
        path = tmp_path / "a.csv"
        path.write_text("1;x\n2;y\n")
        options = CsvByPolarsOptions(
            has_header=False, new_columns=["n", "s"], schema_overrides={"n": pl.String}
        )
        df = data(CsvByPolars(), Csv(), local_item(path), options).collect()
        # The separator still detected; a data type, an object, used.
        assert df.rows() == [("1", "x"), ("2", "y")]
        assert df.columns == ["n", "s"]

    def test_csv_arguments_in_the_code(self, make_item):
        options = CsvByPolarsOptions(
            separator="|", quote_char=None, schema_overrides={"n": pl.String}
        )
        [code] = (
            CsvByPolars()
            .recipe(
                Csv(), make_item("/a.csv", b"x;y\n", location=on_s3("a.csv")), options
            )
            .values()
        )
        # The separator given, not the one detected; None, a value given,
        # shown; data types (objects) not.
        assert last_line(code) == (
            "df = pl.read_csv('s3://b/a.csv', separator='|', quote_char=None)"
        )

    def test_unset_arguments_are_not_passed(self):
        assert CsvByPolarsOptions().arguments() == {}
        assert CsvByPolarsOptions(n_rows=None).arguments() == {"n_rows": None}

    def test_parquet_arguments(self, tmp_path, make_item):
        path = tmp_path / "a.parquet"
        pl.DataFrame({"a": [1, 2, 3]}).write_parquet(path)
        options = ParquetByPolarsOptions(n_rows=2, row_index_name="i")
        df = data(ParquetByPolars(), Parquet(), local_item(path), options).collect()
        assert df.rows() == [(0, 1), (1, 2)]
        # cache is scan_parquet's only: the DataFrame is scanned, then read.
        code = ParquetByPolars().recipe(
            Parquet(),
            make_item("/a.parquet", b"", location=on_s3("a.parquet")),
            ParquetByPolarsOptions(cache=False),
        )
        assert last_line(code["Load as DataFrame"]) == (
            "df = pl.scan_parquet('s3://b/a.parquet', cache=False).collect()"
        )

    def test_excel_sheet(self, make_item):
        item = make_item("/a.xlsx", workbook())
        options = ExcelByPolarsOptions(sheet_name="two")
        assert data(ExcelByPolars(), Excel(), item, options).columns == ["b"]
        [code] = ExcelByPolars().recipe(Excel(), item, options).values()
        assert "df = pl.read_excel(f.read(), sheet_name='two')" in code

    @pytest.mark.parametrize(
        "options",
        [ExcelByPolarsOptions(sheet_id=1), ExcelByPolarsOptions(sheet_name="one")],
    )
    def test_excel_sheet_asked_for(self, make_item, options):
        item = make_item("/a.xlsx", workbook())
        assert data(ExcelByPolars(), Excel(), item, options).columns == ["a"]

    def test_excel_sheets_asked_for_on_tabs(self, make_item):
        item = make_item("/a.xlsx", workbook())
        options = ExcelByPolarsOptions(sheet_id=0)  # all sheets
        view = ExcelByPolars().view(Excel(), item, options, LIMITS)
        assert isinstance(view, TableSetView)
        assert [leaf.id for leaf in view.leaves()] == ["one", "two"]
        assert view.loader("two").columns == ["b"]

    def test_json_code(self, make_item):
        options = JsonByPolarsOptions(infer_schema_length=None)
        [code] = (
            JsonByPolars()
            .recipe(
                Json(), make_item("/a.json", b"[]", location=on_s3("a.json")), options
            )
            .values()
        )
        assert last_line(code) == (
            "df = pl.read_json('s3://b/a.json', infer_schema_length=None)"
        )

    def test_delta_version(self, tmp_path):
        table = tmp_path / "t"
        pl.DataFrame({"a": [1]}).write_delta(table)
        pl.DataFrame({"a": [2]}).write_delta(table, mode="append")
        options = DeltaLakeByPolarsOptions(version=0)
        frame = data(DeltaLakeByPolars(), DeltaLake(), local_item(table), options)
        assert frame.collect()["a"].to_list() == [1]


class TestStoragesPolarsCannotRead:
    """Text formats on a storage polars does not read itself (here memory://)
    are read through the explorer's filesystem; large files once chosen."""

    def test_reads_itself(self, make_item, tmp_path):
        assert reads_itself(local_item(tmp_path))
        assert reads_itself(
            make_item("/a.csv", b"", location=S3Location(bucket="b", key="a.csv"))
        )
        assert not reads_itself(make_item("/a.csv", b""))

    def test_small_files_read_at_once(self, make_item):
        item = make_item("/a.csv", b"a;b\n1;2\n")
        frame = data(CsvByPolars(), Csv(), item).collect()
        assert frame.columns == ["a", "b"]  # the separator detected
        lines = make_item("/a.jsonl", b'{"a": 1}\n{"a": 2}\n')
        assert data(JsonLinesByPolars(), JsonLines(), lines).collect()[
            "a"
        ].to_list() == [1, 2]

    @pytest.mark.parametrize(
        ("label", "rows", "more"),
        [("Its first 1 KB", 1, True), ("The whole file (2 KB)", 2, False)],
    )
    def test_large_files_once_chosen(self, make_item, label, rows, more):
        # Two lines of 1 KB: larger than the files read whole at once.
        line = b"x" * 1023 + b"\n"
        item = make_item("/a.csv", b"name\n" + line + line)
        limits = ViewLimits(max_whole_file_bytes=1100)
        prompt = CsvByPolars().view(Csv(), item, CsvByPolarsOptions(), limits)
        assert isinstance(prompt, ViewWithChoice)
        assert (prompt.style, prompt.with_cancel) == ("button", True)
        assert list(prompt.choices) == ["Its first 1 KB", "The whole file (2 KB)"]
        view = prompt.answer(label)
        assert isinstance(view, QueryView)
        assert (view.data.collect().height, view.more) == (rows, more)

    def test_delta_viewed_only_where_delta_rs_reads(self, make_item, tmp_path):
        make_item("/t/_delta_log/0.json", b"{}")
        support = DeltaLakeByPolars()
        assert not support.supports(
            DeltaLake(), make_item("/t"), DeltaLakeByPolarsOptions(), "view"
        )
        assert support.supports(
            DeltaLake(), local_item(tmp_path), DeltaLakeByPolarsOptions(), "view"
        )

    @pytest.mark.parametrize(
        ("support", "fmt", "name", "reader"),
        [
            (CsvByPolars, Csv, "a.csv", "read_csv(f)"),
            (JsonLinesByPolars, JsonLines, "a.jsonl", "read_ndjson(f)"),
            (JsonByPolars, Json, "a.json", "read_json(f)"),
            (ParquetByPolars, Parquet, "a.parquet", "read_parquet(f)"),
        ],
    )
    def test_code_of_a_file_opened_with_fsspec(
        self, make_item, support, fmt, name, reader
    ):
        # polars cannot read the URL: a DataFrame of the file opened with
        # fsspec, no LazyFrame, which would read it whole all the same.
        item = make_item(f"/{name}", b"a,b\n")
        assert support().supports(fmt(), item, support.options_model(), "recipe")
        options = support.options_model(fsspec_storage_options={"anon": True})
        assert support().recipe(fmt(), item, options) == {
            "Load as DataFrame": (
                "import fsspec\nimport polars as pl\n\n"
                f"with fsspec.open('memory:///{name}', 'rb', anon=True) as f:\n"
                f"    df = pl.{reader}"
            )
        }

    def test_no_code_for_folders_and_delta(self, make_item, memfs):
        memfs.pipe_file("/d/part-0.parquet", b"")
        assert not ParquetByPolars().supports(
            Parquet(), make_item("/d"), ParquetByPolarsOptions(), "recipe"
        )
        memfs.pipe_file("/t/_delta_log/0.json", b"{}")
        assert not DeltaLakeByPolars().supports(
            DeltaLake(), make_item("/t"), DeltaLakeByPolarsOptions(), "recipe"
        )

    def test_excel_code_reads_through_fsspec(self, make_item):
        item = make_item("/a.xlsx", b"")
        assert ExcelByPolars().supports(Excel(), item, ExcelByPolarsOptions(), "recipe")

    def test_parquet_as_a_pyarrow_dataset(self, make_item):
        item = make_item("/a.parquet", parquet_bytes(pl.DataFrame({"a": [1, 2, 3]})))
        frame = data(ParquetByPolars(), Parquet(), item)
        assert isinstance(frame, pl.LazyFrame)  # still lazy
        assert frame.head(2).collect()["a"].to_list() == [1, 2]

    def test_partitioned_parquet_as_a_pyarrow_dataset(self, make_item):
        part = parquet_bytes(pl.DataFrame({"a": [1]}))
        make_item("/ds/year=2026/part-0.parquet", part)
        make_item("/ds/_SUCCESS", b"")  # skipped
        fmt = Parquet()
        item = make_item("/ds")
        assert fmt.matches(item)
        frame = data(ParquetByPolars(), fmt, item).collect()
        assert frame.to_dicts() == [{"a": 1, "year": 2026}]

    def test_parquet_needs_pyarrow_there(self, make_item, tmp_path, monkeypatch):
        monkeypatch.setattr(PyArrow, "is_installed", lambda self: False)
        support = ParquetByPolars()
        assert not support.supports(
            Parquet(), make_item("/a.parquet", b""), ParquetByPolarsOptions(), "view"
        )
        assert support.supports(
            Parquet(), local_item(tmp_path), ParquetByPolarsOptions(), "view"
        )
