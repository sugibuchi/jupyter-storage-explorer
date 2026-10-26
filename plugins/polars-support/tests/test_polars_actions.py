"""polars-support: its actions' views (polars reads by URL: local files) and
code, its options, and its entry point."""

from __future__ import annotations

import io
from importlib.metadata import entry_points

import openpyxl
import polars as pl
import pytest
from jupyter_storage_explorer_polars_support import (
    POLARS_ACTIONS,
    CsvByPolarsOptions,
    DeltaLakeByPolarsOptions,
    ExcelByPolarsOptions,
    JsonByPolarsOptions,
    ParquetByPolarsOptions,
    PolarsOptions,
    PolarsSupport,
    csv_as_data_frame,
    csv_view,
    delta_as_data_frame,
    delta_as_lazy_frame,
    delta_view,
    excel_as_data_frame,
    excel_view,
    json_as_data_frame,
    json_lines_as_data_frame,
    json_lines_as_lazy_frame,
    json_lines_view,
    parquet_as_data_frame,
    parquet_as_lazy_frame,
    parquet_view,
)
from jupyter_storage_explorer_polars_support.actions import _read_by_polars

from jupyter_storage_explorer.filesystems.locations import S3Location
from jupyter_storage_explorer.formats import (
    Csv,
    DeltaLake,
    Excel,
    Json,
    JsonLines,
    Parquet,
)
from jupyter_storage_explorer.packages import PyArrow
from jupyter_storage_explorer.testing import local_item
from jupyter_storage_explorer.views import (
    QueryView,
    TableSetView,
    ViewLimits,
    ViewWithChoice,
)

LIMITS = ViewLimits(max_display_rows=2)


def parquet_bytes(df: pl.DataFrame) -> bytes:
    buffer = io.BytesIO()
    df.write_parquet(buffer)
    return buffer.getvalue()


def workbook() -> bytes:
    """Two sheets: "one", a column "a" of 1, 2, 3; "two", a column "b"."""
    book = openpyxl.Workbook()
    book.active.title = "one"
    for row in (["a"], [1], [2], [3]):
        book.active.append(row)
    book.create_sheet("two").append(["b"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def on_s3(key: str) -> S3Location:
    """A location polars reads itself: its code reads the URL."""
    return S3Location(bucket="b", key=key)


def last_line(code: str) -> str:
    return code.strip().splitlines()[-1]


def data(action, fmt, item, options=None):
    """The data of the view action's view of ``item``: a QueryView's."""
    view = action(fmt, item, options or action.options_model(), LIMITS)
    assert isinstance(view, QueryView)
    return view.data


def test_plugin_registered():
    [entry_point] = entry_points(
        group="jupyter_storage_explorer.plugins", name="polars-support"
    )
    plugin = entry_point.load()()
    assert isinstance(plugin, PolarsSupport)
    assert plugin.description == "Polars support for Storage Explorer"
    assert plugin.actions == POLARS_ACTIONS
    assert csv_view in plugin.actions


def test_one_options_class_per_format():
    # A format's view and code take the same options.
    by_format: dict[type, set[type]] = {}
    for action in POLARS_ACTIONS:
        by_format.setdefault(action.format, set()).add(action.options_model)
    assert all(len(models) == 1 for models in by_format.values())


class TestViews:
    def test_csv_with_its_separator(self, tmp_path):
        path = tmp_path / "a.csv"
        path.write_text("x;y\n1;2\n3;4\n")
        df = data(csv_view, Csv(), local_item(path)).collect()
        assert df.columns == ["x", "y"]  # the separator detected
        assert df.rows() == [(1, 2), (3, 4)]

    @pytest.mark.parametrize(
        ("action", "fmt", "write", "name"),
        [
            (parquet_view, Parquet(), lambda df, p: df.write_parquet(p), "a.parquet"),
            (csv_view, Csv(), lambda df, p: df.write_csv(p), "a.csv"),
            (
                json_lines_view,
                JsonLines(),
                lambda df, p: df.write_ndjson(p),
                "a.jsonl",
            ),
        ],
    )
    def test_local_file(self, tmp_path, action, fmt, write, name):
        path = tmp_path / name
        write(pl.DataFrame({"a": [1, 2]}), path)
        assert data(action, fmt, local_item(path)).collect()["a"].to_list() == [1, 2]

    def test_delta_table(self, tmp_path):
        pl.DataFrame({"a": [1, 2]}).write_delta(tmp_path / "t")
        frame = data(delta_view, DeltaLake(), local_item(tmp_path / "t"))
        assert frame.collect().height == 2

    def test_excel_sheets_on_tabs(self, make_item):
        # polars reads the workbook through the explorer's filesystem: each
        # sheet whole (a DataFrame), queried in memory, when chosen.
        item = make_item("/a.xlsx", workbook())
        view = excel_view(Excel(), item, ExcelByPolarsOptions(), LIMITS)
        assert isinstance(view, TableSetView)
        assert view.selector == "tab"
        assert [(n.label, n.id) for n in view.nodes] == [("one", "one"), ("two", "two")]
        assert view.loader("one")["a"].to_list() == [1, 2, 3]
        assert view.loader("two").columns == ["b"]

    def test_excel_with_one_sheet(self, make_item):
        book = openpyxl.Workbook()
        book.active.append(["a"])
        book.active.append([1])
        buffer = io.BytesIO()
        book.save(buffer)
        item = make_item("/a.xlsx", buffer.getvalue())
        frame = data(excel_view, Excel(), item, ExcelByPolarsOptions())
        assert frame["a"].to_list() == [1]

    def test_parquet_folder_skips_hidden_files(self, tmp_path):
        pl.DataFrame({"a": [1]}).write_parquet(tmp_path / "part-0.parquet")
        pl.DataFrame({"a": [2]}).write_parquet(tmp_path / "part-1.parquet")
        (tmp_path / "_SUCCESS").write_bytes(b"")
        (tmp_path / "sub").mkdir()
        pl.DataFrame({"a": [3]}).write_parquet(tmp_path / "sub" / "part-0.parquet")
        df = data(parquet_view, Parquet(), local_item(tmp_path)).collect()
        assert sorted(df["a"].to_list()) == [1, 2]


class TestPartitionedParquet:
    def write(self, root):
        for year in (2025, 2026):
            for month in (1, 2):
                folder = root / f"year={year}" / f"month={month}"
                folder.mkdir(parents=True)
                pl.DataFrame({"a": [year * 100 + month]}).write_parquet(
                    folder / "part-0.parquet"
                )
                (folder / "_SUCCESS").write_bytes(b"")
        # Left by a writer: not part of the table.
        (root / "_temporary" / "0").mkdir(parents=True)
        pl.DataFrame({"a": [-1]}).write_parquet(root / "_temporary" / "0" / "x.parquet")

    def test_scan_with_partition_columns(self, tmp_path):
        self.write(tmp_path)
        fmt = Parquet()
        assert fmt.matches(local_item(tmp_path))
        df = data(parquet_view, fmt, local_item(tmp_path)).collect().sort("a")
        assert df.columns == ["a", "year", "month"]
        assert df.rows() == [
            (202501, 2025, 1),
            (202502, 2025, 2),
            (202601, 2026, 1),
            (202602, 2026, 2),
        ]

    def test_scan_of_a_folder_not_detected(self, tmp_path):
        # Its layout is found then, through the filesystem given.
        self.write(tmp_path)
        df = data(parquet_view, Parquet(), local_item(tmp_path)).collect()
        assert (df.columns, df.height) == (["a", "year", "month"], 4)

    def test_scan_hive_files(self, tmp_path):
        folder = tmp_path / "dt=2026-10-01"
        folder.mkdir()
        pl.DataFrame({"a": [1, 2]}).write_parquet(folder / "000000_0")
        (folder / "_SUCCESS").write_bytes(b"")
        fmt = Parquet()
        assert fmt.matches(local_item(tmp_path))
        df = data(parquet_view, fmt, local_item(tmp_path)).collect()
        assert df["a"].to_list() == [1, 2]
        assert df.columns == ["a", "dt"]

    def test_code(self, make_item, memfs):
        memfs.pipe_file("/t/year=2025/month=1/part-0.parquet", b"")
        item = make_item("/t")
        options = ParquetByPolarsOptions()
        assert last_line(parquet_as_lazy_frame(Parquet(), item, options)) == (
            "df = pl.scan_parquet('memory:///t/year=*/month=*/*.parquet', "
            "hive_partitioning=True)"
        )
        assert last_line(parquet_as_data_frame(Parquet(), item, options)) == (
            "df = pl.read_parquet('memory:///t/year=*/month=*/*.parquet', "
            "hive_partitioning=True)"
        )

    def test_code_for_hive_files(self, make_item, memfs):
        frame = parquet_bytes(pl.DataFrame({"a": [1]}))
        memfs.pipe_file("/t/dt=2026-10-01/000000_0", frame)
        item = make_item("/t")
        options = ParquetByPolarsOptions()
        arguments = "'memory:///t/dt=*/*', hive_partitioning=True, hidden_file_prefix=['.', '_'])"
        assert last_line(parquet_as_lazy_frame(Parquet(), item, options)) == (
            f"df = pl.scan_parquet({arguments}"
        )
        # read_parquet takes no hidden_file_prefix: scanned, then read.
        assert last_line(parquet_as_data_frame(Parquet(), item, options)) == (
            f"df = pl.scan_parquet({arguments}.collect()"
        )


class TestCode:
    def test_csv(self, make_item):
        code = csv_as_data_frame(
            Csv(),
            make_item("/a.tsv", b"x\ty\n", location=on_s3("a.tsv")),
            CsvByPolarsOptions(),
        )
        assert code.startswith("import polars as pl")
        assert last_line(code) == "df = pl.read_csv('s3://b/a.tsv', separator='\\t')"

    def test_lazy_and_eager(self, make_item, memfs):
        memfs.pipe_file("/d/part-0.parquet", b"")
        folder = make_item("/d")
        options = ParquetByPolarsOptions()
        assert last_line(parquet_as_lazy_frame(Parquet(), folder, options)) == (
            "df = pl.scan_parquet('memory:///d/*.parquet')"
        )
        assert last_line(parquet_as_data_frame(Parquet(), folder, options)) == (
            "df = pl.read_parquet('memory:///d/*.parquet')"
        )
        lines = make_item("/a.jsonl", b"{}\n", location=on_s3("a.jsonl"))
        options = json_lines_view.options_model()
        assert last_line(json_lines_as_lazy_frame(JsonLines(), lines, options)) == (
            "df = pl.scan_ndjson('s3://b/a.jsonl')"
        )
        assert last_line(json_lines_as_data_frame(JsonLines(), lines, options)) == (
            "df = pl.read_ndjson('s3://b/a.jsonl')"
        )

    def test_delta(self, make_item, memfs):
        memfs.pipe_file("/t/_delta_log/0.json", b"{}")
        item = make_item("/t", location=S3Location(bucket="b", key="t"))
        options = DeltaLakeByPolarsOptions()
        assert last_line(delta_as_lazy_frame(DeltaLake(), item, options)) == (
            "df = pl.scan_delta('s3://b/t')"
        )
        assert last_line(delta_as_data_frame(DeltaLake(), item, options)) == (
            "df = pl.read_delta('s3://b/t')"
        )

    def test_json(self, make_item):
        code = json_as_data_frame(
            Json(),
            make_item("/a.json", b"{}", location=on_s3("a.json")),
            JsonByPolarsOptions(),
        )
        assert last_line(code) == "df = pl.read_json('s3://b/a.json')"

    def test_names(self):
        assert [a.label for a in POLARS_ACTIONS if a.format is Parquet] == [
            "View as Parquet (Polars)",
            "Load as LazyFrame",
            "Load as DataFrame",
        ]


class TestOptions:
    def test_storage_options_without_secrets(self, make_item):
        options = CsvByPolarsOptions(
            storage_options={"aws_region": "eu", "aws_secret_access_key": "s"}
        )
        item = make_item("/a.csv", b"x,y\n", location=on_s3("a.csv"))
        assert last_line(csv_as_data_frame(Csv(), item, options)) == (
            "df = pl.read_csv('s3://b/a.csv', storage_options={'aws_region': 'eu'})"
        )

    def test_delta_options(self, make_item, memfs):
        memfs.pipe_file("/t/_delta_log/0.json", b"{}")
        options = DeltaLakeByPolarsOptions(
            storage_options={"region": "eu"},
            use_pyarrow=True,
            delta_table_options={"version": 1},
        )
        item = make_item("/t", location=S3Location(bucket="b", key="t"))
        expected = (
            "'s3://b/t', delta_table_options={'version': 1}, use_pyarrow=True, "
            "storage_options={'region': 'eu'})"
        )
        assert last_line(delta_as_lazy_frame(DeltaLake(), item, options)) == (
            f"df = pl.scan_delta({expected}"
        )
        assert last_line(delta_as_data_frame(DeltaLake(), item, options)) == (
            f"df = pl.read_delta({expected}"
        )

    def test_option_names(self):
        assert PolarsOptions.option_names() == {
            "storage_options",
            "credential_provider",
            "fsspec_storage_options",
        }
        assert DeltaLakeByPolarsOptions.option_names() >= {
            "use_pyarrow",
            "storage_options",
        }


class TestReaderOptions:
    """Each format's options are its reader's arguments, given or left to
    polars."""

    def test_csv_arguments_in_the_data_browser(self, tmp_path):
        path = tmp_path / "a.csv"
        path.write_text("1;x\n2;y\n")
        options = CsvByPolarsOptions(
            has_header=False, new_columns=["n", "s"], schema_overrides={"n": pl.String}
        )
        df = data(csv_view, Csv(), local_item(path), options).collect()
        # The separator still detected; a data type, an object, used.
        assert df.rows() == [("1", "x"), ("2", "y")]
        assert df.columns == ["n", "s"]

    def test_csv_arguments_in_the_code(self, make_item):
        options = CsvByPolarsOptions(
            separator="|", quote_char=None, schema_overrides={"n": pl.String}
        )
        item = make_item("/a.csv", b"x;y\n", location=on_s3("a.csv"))
        # The separator given, not the one detected; None, a value given,
        # shown; data types (objects) not.
        assert last_line(csv_as_data_frame(Csv(), item, options)) == (
            "df = pl.read_csv('s3://b/a.csv', separator='|', quote_char=None)"
        )

    def test_unset_arguments_are_not_passed(self):
        assert CsvByPolarsOptions().arguments() == {}
        assert CsvByPolarsOptions(n_rows=None).arguments() == {"n_rows": None}

    def test_parquet_arguments(self, tmp_path, make_item):
        path = tmp_path / "a.parquet"
        pl.DataFrame({"a": [1, 2, 3]}).write_parquet(path)
        options = ParquetByPolarsOptions(n_rows=2, row_index_name="i")
        df = data(parquet_view, Parquet(), local_item(path), options).collect()
        assert df.rows() == [(0, 1), (1, 2)]
        # cache is scan_parquet's only: the DataFrame is scanned, then read.
        code = parquet_as_data_frame(
            Parquet(),
            make_item("/a.parquet", b"", location=on_s3("a.parquet")),
            ParquetByPolarsOptions(cache=False),
        )
        assert last_line(code) == (
            "df = pl.scan_parquet('s3://b/a.parquet', cache=False).collect()"
        )

    def test_excel_sheet(self, make_item):
        item = make_item("/a.xlsx", workbook())
        options = ExcelByPolarsOptions(sheet_name="two")
        assert data(excel_view, Excel(), item, options).columns == ["b"]
        code = excel_as_data_frame(Excel(), item, options)
        assert "df = pl.read_excel(f.read(), sheet_name='two')" in code

    @pytest.mark.parametrize(
        "options",
        [ExcelByPolarsOptions(sheet_id=1), ExcelByPolarsOptions(sheet_name="one")],
    )
    def test_excel_sheet_asked_for(self, make_item, options):
        item = make_item("/a.xlsx", workbook())
        assert data(excel_view, Excel(), item, options).columns == ["a"]

    def test_excel_sheets_asked_for_on_tabs(self, make_item):
        item = make_item("/a.xlsx", workbook())
        options = ExcelByPolarsOptions(sheet_id=0)  # all sheets
        view = excel_view(Excel(), item, options, LIMITS)
        assert isinstance(view, TableSetView)
        assert [leaf.id for leaf in view.leaves()] == ["one", "two"]
        assert view.loader("two").columns == ["b"]

    def test_json_code(self, make_item):
        options = JsonByPolarsOptions(infer_schema_length=None)
        item = make_item("/a.json", b"[]", location=on_s3("a.json"))
        assert last_line(json_as_data_frame(Json(), item, options)) == (
            "df = pl.read_json('s3://b/a.json', infer_schema_length=None)"
        )

    def test_delta_version(self, tmp_path):
        table = tmp_path / "t"
        pl.DataFrame({"a": [1]}).write_delta(table)
        pl.DataFrame({"a": [2]}).write_delta(table, mode="append")
        options = DeltaLakeByPolarsOptions(version=0)
        frame = data(delta_view, DeltaLake(), local_item(table), options)
        assert frame.collect()["a"].to_list() == [1]


class TestStoragesPolarsCannotRead:
    """Text formats on a storage polars does not read itself (here memory://)
    are read through the explorer's filesystem; large files once chosen."""

    def test_read_by_polars(self, make_item, tmp_path):
        def read(item):
            return _read_by_polars(Csv(), item, CsvByPolarsOptions())

        assert read(local_item(tmp_path))
        assert read(
            make_item("/a.csv", b"", location=S3Location(bucket="b", key="a.csv"))
        )
        assert not read(make_item("/a.csv", b""))

    def test_small_files_read_at_once(self, make_item):
        item = make_item("/a.csv", b"a;b\n1;2\n")
        frame = data(csv_view, Csv(), item).collect()
        assert frame.columns == ["a", "b"]  # the separator detected
        lines = make_item("/a.jsonl", b'{"a": 1}\n{"a": 2}\n')
        frame = data(json_lines_view, JsonLines(), lines).collect()
        assert frame["a"].to_list() == [1, 2]

    @pytest.mark.parametrize(
        ("label", "rows", "more"),
        [("Its first 1.1 KB", 1, True), ("The whole file (2 KB)", 2, False)],
    )
    def test_large_files_once_chosen(self, make_item, label, rows, more):
        # Two lines of 1 KB: larger than the files read whole at once.
        line = b"x" * 1023 + b"\n"
        item = make_item("/a.csv", b"name\n" + line + line)
        limits = ViewLimits(max_whole_file_bytes=1100)
        prompt = csv_view(Csv(), item, CsvByPolarsOptions(), limits)
        assert isinstance(prompt, ViewWithChoice)
        assert (prompt.style, prompt.with_cancel) == ("button", True)
        assert list(prompt.choices) == ["Its first 1.1 KB", "The whole file (2 KB)"]
        view = prompt.answer(label)
        assert isinstance(view, QueryView)
        assert (view.data.collect().height, view.more) == (rows, more)

    def test_delta_viewed_only_where_delta_rs_reads(self, make_item, tmp_path):
        make_item("/t/_delta_log/0.json", b"{}")
        options = DeltaLakeByPolarsOptions()
        assert not delta_view.does(DeltaLake(), make_item("/t"), options)
        assert delta_view.does(DeltaLake(), local_item(tmp_path), options)

    @pytest.mark.parametrize(
        ("action", "fmt", "name", "reader"),
        [
            (csv_as_data_frame, Csv, "a.csv", "read_csv(f)"),
            (json_lines_as_data_frame, JsonLines, "a.jsonl", "read_ndjson(f)"),
            (json_as_data_frame, Json, "a.json", "read_json(f)"),
            (parquet_as_data_frame, Parquet, "a.parquet", "read_parquet(f)"),
        ],
    )
    def test_code_of_a_file_opened_with_fsspec(
        self, make_item, action, fmt, name, reader
    ):
        # polars cannot read the URL: a DataFrame of the file opened with
        # fsspec, no LazyFrame, which would read it whole all the same.
        item = make_item(f"/{name}", b"a,b\n")
        assert action.does(fmt(), item, action.options_model())
        lazy = [
            a
            for a in POLARS_ACTIONS
            if a.format is fmt and a.label == "Load as LazyFrame"
        ]
        assert not any(a.does(fmt(), item, a.options_model()) for a in lazy)
        options = action.options_model(fsspec_storage_options={"anon": True})
        assert action(fmt(), item, options) == (
            "import fsspec\nimport polars as pl\n\n"
            f"with fsspec.open('memory:///{name}', 'rb', anon=True) as f:\n"
            f"    df = pl.{reader}"
        )

    def test_no_code_for_folders_and_delta(self, make_item, memfs):
        memfs.pipe_file("/d/part-0.parquet", b"")
        options = ParquetByPolarsOptions()
        for action in (parquet_as_lazy_frame, parquet_as_data_frame):
            assert not action.does(Parquet(), make_item("/d"), options)
        memfs.pipe_file("/t/_delta_log/0.json", b"{}")
        options = DeltaLakeByPolarsOptions()
        for action in (delta_as_lazy_frame, delta_as_data_frame):
            assert not action.does(DeltaLake(), make_item("/t"), options)

    def test_excel_code_reads_through_fsspec(self, make_item):
        item = make_item("/a.xlsx", b"")
        assert excel_as_data_frame.does(Excel(), item, ExcelByPolarsOptions())

    def test_parquet_as_a_pyarrow_dataset(self, make_item):
        item = make_item("/a.parquet", parquet_bytes(pl.DataFrame({"a": [1, 2, 3]})))
        frame = data(parquet_view, Parquet(), item)
        assert isinstance(frame, pl.LazyFrame)  # still lazy
        assert frame.head(2).collect()["a"].to_list() == [1, 2]

    def test_partitioned_parquet_as_a_pyarrow_dataset(self, make_item):
        part = parquet_bytes(pl.DataFrame({"a": [1]}))
        make_item("/ds/year=2026/part-0.parquet", part)
        make_item("/ds/_SUCCESS", b"")  # skipped
        fmt = Parquet()
        item = make_item("/ds")
        assert fmt.matches(item)
        frame = data(parquet_view, fmt, item).collect()
        assert frame.to_dicts() == [{"a": 1, "year": 2026}]

    def test_parquet_needs_pyarrow_there(self, make_item, tmp_path, monkeypatch):
        monkeypatch.setattr(PyArrow, "is_installed", lambda self: False)
        options = ParquetByPolarsOptions()
        assert not parquet_view.does(Parquet(), make_item("/a.parquet", b""), options)
        assert parquet_view.does(Parquet(), local_item(tmp_path), options)
